"""Golden tests on Huggies FreeFeel (N=406) - spec section 11.1.  Needs the SAV mounted (never in the repo)."""
import io

import numpy as np
import openpyxl

from conftest import needs_golden

pytestmark = needs_golden


def v(g, united, level="sample"):
    return g["vals"][united][level][0]


def test_levels(golden):
    b = golden["ctx"].bases()
    assert b["sample"] == 406
    assert b["exposed"] == 217
    assert b["customers"] == 246 and b["noncust"] == 160
    assert b["notexposed"] == 189


def test_exposure_union_equals_dp_variable(golden):
    chk = [c for c in golden["ctx"].checks if "vctq1c1" in c[0]]
    assert chk and chk[0][1].startswith("406/406") and chk[0][2] == "OK"


def test_follow_up_bases(golden):
    assert round(v(golden, "SEMIAEX-VD#01"), 1) == 29.6
    assert round(v(golden, "SEMIAEX-CAT#01"), 1) == 37.4
    assert round(v(golden, "UNEXPOSEB#01"), 1) == 36.0
    assert round(v(golden, "SLOGAN#01"), 1) == 24.9


def test_attention_range(golden):
    assert 0 <= v(golden, "DESCRIBE#27") <= 42


def test_single_choice_sums_to_100(golden):
    bad = [c for c in golden["ctx"].checks if "סכום התשובות = 100" in c[0] and c[2] != "OK"]
    assert not bad, bad


def test_no_engine_errors_or_flags(golden):
    assert golden["summary"]["errors"] == []
    assert golden["summary"]["flagged"] == []


def test_workbook_structure_and_rtl(golden):
    wb = openpyxl.load_workbook(io.BytesIO(golden["xlsx"]))
    assert wb.sheetnames == ["README", "ממצאים", "משתנים מסכמים", "DATA_לייבוא", "הגדרות", "בקרות", "לא חושב"]
    assert all(ws.sheet_view.rightToLeft for ws in wb.worksheets)
    readme = "\n".join(str(r[0]) for r in wb["README"].iter_rows(values_only=True) if r[0])
    assert "גרסת אפליקציה" in readme and "build" in readme


def test_no_formula_strings(golden):
    wb = openpyxl.load_workbook(io.BytesIO(golden["xlsx"]))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert c.data_type == "s", (ws.title, c.coordinate)


def test_customer_cuts_never_exported(golden):
    wb = openpyxl.load_workbook(io.BytesIO(golden["xlsx"]))
    bases = {r[6] for r in wb["DATA_לייבוא"].iter_rows(min_row=2, values_only=True) if r[4]}
    assert bases <= {"מדגם", "נחשפים"}


def test_deterministic_values(golden):
    from svr import pipeline as P
    from conftest import SAV, DICT
    _, _, ctx2 = P.run_findings(str(SAV), golden["mapping"], str(DICT))
    v2 = {r["united"]: r["vals"] for t in ctx2.tables for r in t["rows"] if r["united"]}
    assert v2 == golden["vals"]
