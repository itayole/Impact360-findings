"""End-to-end API tests on the Huggies example (spec 11.3): upload -> review -> approve -> run -> download."""
import importlib
import io
import json
import os
import time

import openpyxl
import pytest
from fastapi.testclient import TestClient

from conftest import needs_golden, ROOT, SAV

DOCX = ROOT / "Example" / "שאלון אפקטיביות קמפיין השקת האגיס  FREE FEEL final (ID 146445).docx"
CONF = json.loads((ROOT / "tests" / "golden" / "haggies_confirmations.json").read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    os.environ["DATA_DIR"] = str(tmp_path_factory.mktemp("data"))
    from app import config, store, runtime, service, main
    for m in (config, store, runtime, service, main):
        importlib.reload(m)
    return TestClient(main.app)


def wait_job(c, jid, timeout=120):
    t0 = time.time()
    while time.time() - t0 < timeout:
        j = c.get(f"/api/jobs/{jid}").json()
        if j["status"] in ("done", "error"):
            return j
        time.sleep(0.5)
    raise AssertionError("job timeout")


def test_health_version(client):
    assert client.get("/api/health").json()["status"] == "ok"
    v = client.get("/api/version").json()
    assert v["app_version"] and "145212" in v["dictionary"]


@needs_golden
def test_full_flow(client):
    files = {"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}
    if DOCX.exists():
        files["qnr"] = (DOCX.name, open(DOCX, "rb"), "application/octet-stream")
    r = client.post("/api/projects", files=files, data={"name": "האגיס בדיקה"}, headers={"x-user": "tester"})
    assert r.status_code == 200, r.text
    d = r.json()
    pid = d["project"]["id"]
    assert "האגיס" in d["brand_options"] and d["project"]["n"] == 406

    r = client.post(f"/api/projects/{pid}/profile", json={"brand": "האגיס", "campaign_id": "TEST-1"})
    assert r.status_code == 200, r.text
    rv = r.json()
    assert set(rv["pending"]) >= {"q31", "seen"} and rv["levels"]["bases"]["exposed"] == 217
    assert rv["messages"]["options"] and [o["role"] for o in rv["messages"]["options"]][:4] == ["main", "sec1", "sec2", "buy"]

    # approvals + Net Builder decisions (the Huggies golden decisions)
    patch = {"questions": {k: {"confirmed": True} for k in CONF["confirmed"]}}
    for k, nets in CONF["nets"].items():
        patch["questions"].setdefault(k, {})["nets"] = nets
    patch["questions"]["q31"]["nets"] = CONF["q31_nets"]
    r = client.put(f"/api/projects/{pid}/mapping", json=patch, headers={"x-user": "tester"})
    assert r.status_code == 200, r.text

    # live preview == what the workbook will contain
    pv = client.post(f"/api/projects/{pid}/preview", json={"block": "vSLOGAN_coded"}).json()
    slog = next(n for n in pv["nets"] if n["united"] == "SLOGAN#01")
    assert round(slog["values"]["sample"], 1) == 24.9
    assert all(c["counts"] is not None for c in pv["codes"])
    assert slog["counts"]["sample"] == round(24.9 * 406 / 100)      # net rows carry n like the answer rows

    j = client.post(f"/api/projects/{pid}/run").json()
    res = wait_job(client, j["job_id"])
    assert res["status"] == "done", res
    assert res["result"]["bases"]["exposed"] == 217 and res["result"]["summary"]["errors"] == []

    rs = client.get(f"/api/projects/{pid}/results").json()
    row = next(r for t in rs["tables"] for r in t["rows"] if r["united"] == "SLOGAN#01")
    assert round(row["values"]["sample"], 1) == 24.9 and row["counts"]["sample"] == round(24.9 * 406 / 100)
    assert set(rs["letters"].values()) == {"A", "B", "C", "D", "E"}

    x = client.get(f"/api/projects/{pid}/download")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    data = {r[4]: r[3] for r in wb["DATA_לייבוא"].iter_rows(min_row=2, values_only=True) if r[4]}
    assert data["SEMIAEX-CAT#01"] == 37.4 and data["UNEXPOSEB#01"] == 36.0 and data["SEMIAEX-VD#01"] == 29.6
    # decision log lands in the workbook (who / when / what)
    log_rows = [r for r in wb["הגדרות"].iter_rows(values_only=True) if r[1] == "tester"]
    assert log_rows

    # every open match must be approved or excluded before a template can be saved
    assert client.post("/api/library/templates", json={"project_id": pid, "name": "Huggies FF"}).status_code == 400
    client.put(f"/api/projects/{pid}/mapping", json={"questions": {"seen": {"include": False}}}, headers={"x-user": "tester"})
    # template round trip through the API
    r = client.post("/api/library/templates", json={"project_id": pid, "name": "Huggies FF", "client": "K", "tracker": "FF"})
    assert r.status_code == 200, r.text
    r2 = client.post("/api/projects", files={"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}, data={"name": "גל 2"})
    assert r2.json()["template_suggestions"][0]["similarity"] == 1.0
    pid2 = r2.json()["project"]["id"]
    rv2 = client.post(f"/api/projects/{pid2}/profile", json={"brand": "האגיס", "template_id": "Huggies_FF"}).json()
    assert rv2["pending"] == []                             # wave 2: zero manual approvals (spec 11.3: < 5)
    assert rv2["mapping"]["template_diff"]["missing"] == []


@needs_golden
def test_remap_and_warnings(client):
    pid = client.post("/api/projects", files={"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}).json()["project"]["id"]
    client.post(f"/api/projects/{pid}/profile", json={"brand": "האגיס"})
    r = client.put(f"/api/projects/{pid}/mapping", json={"remap": {"q54": "ENJOY"}}, headers={"x-user": "dp1"})
    assert r.status_code == 200, r.text
    m = {e["key"]: e for e in r.json()["mapping"]["questions"]}
    assert m["q54"]["dict_var"] == "ENJOY" and m["q54"]["confidence"] == "manual" and m["ENJOY"]["dict_var"] is None
    assert any(d["user"] == "dp1" for d in client.get(f"/api/projects/{pid}/decisions").json())
    w = client.get(f"/api/projects/{pid}/warnings").json()
    assert any(x["kind"] == "unconfirmed" for x in w)
    bad = client.put(f"/api/projects/{pid}/mapping", json={"questions": {"q54": {"dict_var": "X"}}})
    assert bad.status_code == 400
    # a template with unapproved matches cannot be saved
    assert client.post("/api/library/templates", json={"project_id": pid, "name": "x"}).status_code == 400


@needs_golden
def test_message_roles(client):
    pid = client.post("/api/projects", files={"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}).json()["project"]["id"]
    rv = client.post(f"/api/projects/{pid}/profile", json={"brand": "האגיס"}).json()
    n = len(rv["messages"]["options"])
    roles = ["buy", "main", "sec1", "sec2"][:n]
    r = client.put(f"/api/projects/{pid}/mapping", json={"roles": roles})
    assert r.status_code == 200, r.text
    assert [o["role"] for o in r.json()["messages"]["options"]] == roles
    assert client.put(f"/api/projects/{pid}/mapping", json={"roles": ["main", "main", "", ""][:n]}).status_code == 400


@needs_golden
def test_snapshots_save_load_delete(client):
    pid = client.post("/api/projects", files={"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}).json()["project"]["id"]
    client.post(f"/api/projects/{pid}/profile", json={"brand": "האגיס"})
    client.put(f"/api/projects/{pid}/mapping", json={"questions": {"q33": {"confirmed": True}}}, headers={"x-user": "u1"})
    sid = client.post(f"/api/projects/{pid}/snapshots", json={"name": "לפני שינוי"}, headers={"x-user": "u1"}).json()["id"]
    client.put(f"/api/projects/{pid}/mapping", json={"questions": {"q33": {"confirmed": False}, "q49": {"include": False}}})
    rv = client.post(f"/api/projects/{pid}/snapshots/{sid}/load", json={}).json()
    q = {e["key"]: e for e in rv["mapping"]["questions"]}
    assert q["q33"]["confirmed"] is True and q["q49"].get("include", True) is True
    names = [s["name"] for s in client.get(f"/api/projects/{pid}/snapshots").json()]
    assert "לפני שינוי" in names and any(n.startswith("גיבוי אוטומטי") for n in names)
    assert client.delete(f"/api/projects/{pid}/snapshots/{sid}").status_code == 200
    assert client.post(f"/api/projects/{pid}/snapshots/{sid}/load", json={}).status_code == 404
    assert client.post(f"/api/projects/{pid}/snapshots", json={"name": " "}).status_code == 400


@needs_golden
def test_custom_customer_definition(client):
    pid = client.post("/api/projects", files={"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}).json()["project"]["id"]
    rv = client.post(f"/api/projects/{pid}/profile", json={"brand": "האגיס"}).json()
    assert rv["levels"]["bases"]["customers"] == 246
    names = [v["var"] for v in client.get(f"/api/projects/{pid}/variables").json()]
    assert "USAGEr1" in names and "USAGEr2" in names
    d = client.get(f"/api/projects/{pid}/variables/USAGEr2").json()
    assert d["n"] == 406 and sum(v["n"] for v in d["values"]) + d["n_missing"] == 406
    r = client.put(f"/api/projects/{pid}/mapping", json={"customer_def": {"var": "USAGEr2", "values": [1]}}, headers={"x-user": "u"})
    assert r.status_code == 200, r.text
    b = r.json()["levels"]["bases"]
    assert b["customers"] == next(v["n"] for v in d["values"] if v["value"] == 1) and b["customers"] + b["noncust"] == 406
    assert any("USAGEr2" in x["what"] for x in client.get(f"/api/projects/{pid}/decisions").json())
    assert client.put(f"/api/projects/{pid}/mapping", json={"customer_def": {"var": "USAGEr2", "values": []}}).status_code == 400
    assert client.put(f"/api/projects/{pid}/mapping", json={"customer_def": {"var": "nope", "values": [1]}}).status_code == 400
    assert client.put(f"/api/projects/{pid}/mapping", json={"customer_def": {"var": "USAGEr2", "values": [7]}}).status_code == 400


@needs_golden
def test_project_name_follows_the_form_and_renames(client):
    pid = client.post("/api/projects", files={"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}, data={"name": "ישן"}).json()["project"]["id"]
    rv = client.post(f"/api/projects/{pid}/profile", json={"brand": "האגיס", "name": "שם סופי"}).json()
    assert rv["mapping"]["project"]["name"] == "שם סופי"
    assert next(p for p in client.get("/api/projects").json() if p["id"] == pid)["name"] == "שם סופי"
    client.put(f"/api/projects/{pid}/mapping", json={"project": {"name": "שם חדש"}})
    assert next(p for p in client.get("/api/projects").json() if p["id"] == pid)["name"] == "שם חדש"


@needs_golden
def test_summary_order_is_saved_and_drives_row_order(client):
    pid = client.post("/api/projects", files={"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}).json()["project"]["id"]
    client.post(f"/api/projects/{pid}/profile", json={"brand": "האגיס"})
    a = {"SLOGAN#01": {"include": ["vSLOGAN_codedr2"]}, "USER:x": {"include": ["vSLOGAN_codedr3"], "label": "שני"}}
    client.put(f"/api/projects/{pid}/mapping", json={"questions": {"vSLOGAN_coded": {"nets": a}}})
    b = {"USER:x": a["USER:x"], "SLOGAN#01": a["SLOGAN#01"]}
    client.put(f"/api/projects/{pid}/mapping", json={"questions": {"vSLOGAN_coded": {"nets": b}}})      # order only
    cat = next(x for x in client.get(f"/api/projects/{pid}/catalog").json() if x["key"] == "vSLOGAN_coded")
    assert list(cat["nets"]) == ["USER:x", "SLOGAN#01"]
    pv = client.post(f"/api/projects/{pid}/preview", json={"block": "vSLOGAN_coded"}).json()
    assert [n["label"] for n in pv["nets"]][0] == "שני"
    client.delete(f"/api/projects/{pid}")


@needs_golden
def test_summary_members_in_results_and_workbook(client):
    pid = client.post("/api/projects", files={"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}).json()["project"]["id"]
    client.post(f"/api/projects/{pid}/profile", json={"brand": "האגיס"})
    nets = {"SLOGAN#01": {"include": ["vSLOGAN_codedr2"]}, "USER:x": {"exclude": ["vSLOGAN_codedr1"], "label": "הכול חוץ מלא יודע"}}
    client.put(f"/api/projects/{pid}/mapping", json={"questions": {"vSLOGAN_coded": {"nets": nets, "confirmed": True}}})
    rs = client.get(f"/api/projects/{pid}/results").json()
    tb = next(t for t in rs["tables"] if t["key"] == "vSLOGAN_coded")
    inc = next(r for r in tb["rows"] if r["united"] == "SLOGAN#01")
    exc = next(r for r in tb["rows"] if r["label"] == "הכול חוץ מלא יודע")
    assert inc["members_text"] == "כולל: We got you, baby" and exc["members_text"].startswith("כל התשובות מלבד: ")
    tagged = [r for r in tb["rows"] if r.get("in_nets")]
    assert len(tagged) == 1 and tagged[0]["in_nets"] == ["סה\"כ ציינו סיסמא נכונה"]
    j = client.post(f"/api/projects/{pid}/run").json()["job_id"]
    res = wait_job(client, j)
    assert res["status"] == "done"
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/projects/{pid}/download").content))
    notes = [r[-1] for r in wb["ממצאים"].iter_rows(values_only=True) if r and r[1] == "SLOGAN#01"]
    assert any("כולל: We got you, baby" in (n or "") for n in notes)
    data_notes = [r[8] for r in wb["DATA_לייבוא"].iter_rows(min_row=2, values_only=True) if r[4] == "SLOGAN#01"]
    assert all("כולל" not in (n or "") for n in data_notes)             # DATA_לייבוא stays as it was
    client.delete(f"/api/projects/{pid}")


@needs_golden
def test_columns_exposed_nonusers_and_segments(client):
    pid = client.post("/api/projects", files={"sav": (SAV.name, open(SAV, "rb"), "application/octet-stream")}).json()["project"]["id"]
    rv = client.post(f"/api/projects/{pid}/profile", json={"brand": "האגיס"}).json()
    cols = rv["levels"]["columns"]
    assert [c["key"] for c in cols] == ["sample", "exposed", "expnonuser", "customers", "noncust"]
    assert [c["letter"] for c in cols] == list("ABCDE") and cols[2]["name"] == "נחשפים שאינם משתמשים" and cols[2]["n"] == 74
    r = client.put(f"/api/projects/{pid}/mapping", json={"segment_add": {"name": "קונות פמפרס", "var": "USAGEr3", "values": [1]}}, headers={"x-user": "u"})
    assert r.status_code == 200, r.text
    cols = r.json()["levels"]["columns"]
    assert cols[-1]["letter"] == "F" and cols[-1]["segment"] and cols[-1]["name"] == "קונות פמפרס"
    rs = client.get(f"/api/projects/{pid}/results").json()
    row = next(r_ for t in rs["tables"] for r_ in t["rows"] if r_["united"] == "ENJOY#91")
    key = cols[-1]["key"]
    assert row["values"][key] is not None and row["n"][key] == cols[-1]["n"] and row["letters"]["customers"] in ("", "E", "e")
    assert "notexposed" not in row["values"] and row["letters"]["exposed"] == "" and row["letters"]["expnonuser"] == ""
    # analysis only: the new columns never reach DATA_import
    j = client.post(f"/api/projects/{pid}/run").json()["job_id"]
    assert wait_job(client, j)["status"] == "done"
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/projects/{pid}/download").content))
    assert {r_[6] for r_ in wb["DATA_לייבוא"].iter_rows(min_row=2, values_only=True) if r_[4]} <= {"מדגם", "נחשפים"}
    hdr = [c.value for c in wb["ממצאים"][5]] if False else None
    assert any("קונות פמפרס" in str(c.value) for row_ in wb["ממצאים"].iter_rows(max_row=12) for c in row_)
    assert client.put(f"/api/projects/{pid}/mapping", json={"segment_add": {"name": "", "var": "USAGEr3", "values": [1]}}).status_code == 400
    assert client.put(f"/api/projects/{pid}/mapping", json={"segment_add": {"name": "x", "var": "USAGEr3", "values": [9]}}).status_code == 400
    sid = r.json()["mapping"]["segments"][0]["id"]
    r2 = client.put(f"/api/projects/{pid}/mapping", json={"segment_remove": sid})
    assert r2.status_code == 200 and len(r2.json()["levels"]["columns"]) == 5
    client.delete(f"/api/projects/{pid}")


def test_upload_validation(client):
    r = client.post("/api/projects", files={"sav": ("x.txt", b"abc", "text/plain")})
    assert r.status_code == 400
    r = client.post("/api/projects", files={"sav": ("x.sav", b"not spss", "application/octet-stream")})
    assert r.status_code == 400 and "SAV" in r.json()["detail"]


def test_retention_cleanup(client):
    from app import main
    st = main.store
    old = os.path.join(st.uploads, "abc.sav")
    open(old, "wb").write(b"x")
    past = time.time() - 40 * 86400
    os.utime(old, (past, past))
    assert st.cleanup()["uploads"] >= 1 and not os.path.exists(old)


def test_gone_after_cleanup(client):
    from app import main
    p = main.store.new_project("t", "t.sav", "x.xlsx")
    r = client.post(f"/api/projects/{p['id']}/profile", json={"brand": ""})
    assert r.status_code == 410
