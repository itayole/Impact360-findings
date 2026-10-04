"""Regression tests for the QA findings (charts correctness, template paths, design-file handling, uploads, locks)."""
import io
import os
import threading
import zipfile

import pytest
from pptx import Presentation
from pptx.util import Inches

from svr import charts as CH
from svr import templates as T


def _row(label, v, n=500, section="item", role="", kind="pct"):
    return dict(label=label, united="X#01", section=section, role=role, kind=kind, note="", values={"sample": v}, n={"sample": n}, letters={}, counts={})


def _res(rows, typ="single", **kw):
    return dict(columns=[], tables=[dict(key="K", title="K", question="q", dict_var="K", type=typ, rows=rows, **kw)])


# ------------------------------------------------------------------------------------------ charts
def test_chart_base_comes_from_the_answers_and_a_summary_on_another_base_is_left_out():
    rows = [_row("תשובה א", 60.0, n=406), _row("תשובה ב", 40.0, n=406), _row("סיכום על נחשפים", 53.0, n=217, section="summary", role="item")]
    s = CH.build_specs(_res(rows))[0]
    assert s["base_n"] == 406 and s["dropped_head"] == 1
    assert [c["headline"] for c in s["categories"]] == [False, False]
    prs = Presentation(io.BytesIO(CH.render_pptx([s])))
    assert any("סיכום על בסיס אחר לא מוצג" in sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    same = CH.build_specs(_res([_row("א", 60.0, n=406), _row("סיכום", 60.0, n=406, section="summary", role="item")]))[0]
    assert same["dropped_head"] == 0 and same["categories"][0]["headline"]


def test_control_characters_and_non_finite_values_do_not_break_the_export():
    rows = [_row("תווית\x01פגומה", 40.0), _row("ריקה", float("nan")), _row("אינסוף", float("inf")), _row("תקינה", 60.0)]
    s = CH.build_specs(_res(rows))[0]
    assert [c["label"] for c in s["categories"]] == ["תוויתפגומה", "תקינה"]
    for ct in CH.CHART_TYPES:
        s["chart_type"] = ct
        Presentation(io.BytesIO(CH.render_pptx([s])))               # must not raise


def test_bad_settings_are_dropped_not_500():
    for raw in ("x", [1], dict(defaults="x"), dict(defaults=dict(chart_type=[1], color=5)), dict(questions=[1]),
                dict(questions=dict(a=dict(chart_type=[1], color="#12345\n"))), dict(defaults=dict(color="#112233\n"))):
        st = CH.clean_settings(raw)
        assert st["defaults"]["chart_type"] == "bar_h" and st["defaults"]["color"] == CH.DEFAULT_COLOR and st["questions"] == {}, raw


def test_pie_and_stacked_only_for_answers_that_add_up():
    rows = [_row("א", 60.0), _row("ב", 40.0)]
    for typ, ok in (("single", True), ("scale", True), ("multi", False), ("coded_open", False), ("derived", False)):
        s = CH.build_specs(_res(rows, typ), dict(defaults=dict(chart_type="donut")))[0]
        assert (s["chart_type"] == "donut") is ok, typ
        assert ("donut" in s["allowed_types"]) is ok, typ


def test_axis_leaves_room_for_a_100_percent_label_and_is_never_degenerate():
    for v in (100.0, 0.0):
        s = CH.build_specs(_res([_row("א", v)]))[0]
        prs = Presentation(io.BytesIO(CH.render_pptx([s])))
        ch = next(sh.chart for sh in prs.slides[0].shapes if sh.has_chart)
        assert ch.value_axis.maximum_scale >= (120 if v else 10)


def test_footer_drops_piping_tokens_and_is_cut_to_two_lines():
    spec = dict(title="החיתול הכי טוב של המותג - [pipe: TOTAL_MESSAGEtxt]", key="K", dict_var="K")
    assert "[pipe" not in CH.question_footer(spec)
    long = CH.question_footer(dict(title="ש" * 600, key="K", dict_var="K"))
    assert len(long) < 275 and "…" in long


def test_chart_text_is_rtl_inside_the_chart():
    s = CH.build_specs(_res([_row("האגיס (כללי)", 60.0), _row("ב", 40.0)]))[0]
    z = zipfile.ZipFile(io.BytesIO(CH.render_pptx([s])))
    x = z.read("ppt/charts/chart1.xml").decode("utf-8")
    assert 'rtl="1"' in x


def _base_with_sections():
    prs = Presentation()
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[5])
    buf = io.BytesIO()
    prs.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    ids = [int(i) for i in __import__("re").findall(r'<p:sldId id="(\d+)"', zin.read("ppt/presentation.xml").decode())]
    ext = ('<p:extLst><p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"><p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
           '<p14:section name="S" id="{11111111-1111-1111-1111-111111111111}"><p14:sldIdLst>' + "".join(f'<p14:sldId id="{i}"/>' for i in ids) +
           '</p14:sldIdLst></p14:section></p14:sectionLst></p:ext></p:extLst>')
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zin.infolist():
            b = zin.read(it.filename)
            if it.filename == "ppt/presentation.xml":
                b = b.decode("utf-8").replace("</p:presentation>", ext + "</p:presentation>").encode("utf-8")
            zo.writestr(it, b)
    return out.getvalue()


def test_client_deck_sections_do_not_point_at_deleted_slides():
    base = _base_with_sections()
    x0 = zipfile.ZipFile(io.BytesIO(base)).read("ppt/presentation.xml").decode()
    assert "sectionLst" in x0
    data = CH.render_pptx(CH.build_specs(_res([_row("א", 60.0)])), base=base)
    x = zipfile.ZipFile(io.BytesIO(data)).read("ppt/presentation.xml").decode()
    assert "sectionLst" not in x and "custShowLst" not in x
    assert len(Presentation(io.BytesIO(data)).slides) == 1


# ------------------------------------------------------------------------------------------ templates
@pytest.mark.parametrize("tid", ["../evil", "..\\evil", "C:/Windows/Temp/x", "a/b", "", "x" * 61, "a\n", None, 5])
def test_template_ids_cannot_leave_the_library(tmp_path, tid):
    lib = T.Library(str(tmp_path / "lib"))
    assert not T.valid_tid(tid)
    assert lib.get(tid) is None and lib._versions(tid) == [] and lib.find_base(tid, 1) is None
    if isinstance(tid, str) and tid:
        with pytest.raises(ValueError):
            lib.save("n", dict(project={}, questions=[]), [], template_id=tid)
    assert not (tmp_path / "evil").exists()


def test_template_name_collision_is_refused_but_a_new_version_of_the_same_name_is_fine(tmp_path):
    lib = T.Library(str(tmp_path / "lib"))
    m = dict(project={}, questions=[])
    assert lib.save("A B", m, [])["version"] == 1
    assert lib.save("A B", m, [])["version"] == 2                   # same name = same template
    with pytest.raises(ValueError):
        lib.save("A_B", m, [])                                      # different name, same folder id
    assert lib.save("A_B", m, [], template_id="A_B")["version"] == 3


def test_design_file_is_stored_per_template_version(tmp_path):
    lib = T.Library(str(tmp_path / "lib"))
    lib.save("T", dict(project={}, questions=[]), [])
    p = lib.base_path("T", 1)
    open(p, "wb").write(b"x")
    assert lib.find_base("T", 1) == p and lib.find_base("T", 2) is None
    open(os.path.join(lib.tdir, "T", "charts_base.pptx"), "wb").write(b"legacy")        # written by 0.3.0 / 0.3.1
    assert lib.find_base("T", 2).endswith("charts_base.pptx")


# ------------------------------------------------------------------------------------------ store
def test_concurrent_writes_never_collide_and_a_deleted_project_is_not_resurrected(tmp_path):
    from app import store as S
    st = S.Store(str(tmp_path))
    p = st.new_project("n", "a.sav", "d.xlsx")
    errs = []

    def w():
        try:
            for _ in range(40):
                q = st.project(p["id"])
                st.save_project(q)
        except Exception as e:  # noqa: BLE001
            errs.append(repr(e))
    ts = [threading.Thread(target=w) for _ in range(6)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert errs == [] and st.project(p["id"])["id"] == p["id"]
    st.delete_project(p["id"])
    with pytest.raises(KeyError):
        st.save_project(p)                                          # a job finishing after the delete
    with pytest.raises(KeyError):
        st.save_json(p["id"], "x.json", {})
    assert not os.path.exists(os.path.join(st.projects, p["id"]))
    assert not S.PID_RE.match(p["id"] + "\n") and not S.SID_RE.match("abcdef12\n")


def test_cleanup_survives_a_file_it_cannot_remove_and_purges_orphan_folders(tmp_path, monkeypatch):
    from app import store as S
    st = S.Store(str(tmp_path))
    open(os.path.join(st.uploads, "a.sav"), "wb").write(b"1")
    os.makedirs(os.path.join(st.uploads, "subdir"))
    orphan = os.path.join(st.projects, "0123456789ab")
    os.makedirs(orphan)                                              # no project.json
    real = os.remove
    monkeypatch.setattr(S.os, "remove", lambda p: (_ for _ in ()).throw(PermissionError("locked")) if p.endswith("a.sav") else real(p))
    r = st.cleanup(now_ts=10 ** 10)                                  # everything is "old"
    assert r["projects"] == 1 and not os.path.exists(orphan)


# ------------------------------------------------------------------------------------------ API
@pytest.fixture()
def api(tmp_path):
    import importlib
    os.environ["DATA_DIR"] = str(tmp_path)
    from app import config, store, runtime, service, main
    for m in (config, store, runtime, service, main):
        importlib.reload(m)
    from fastapi.testclient import TestClient
    return TestClient(main.app), main, service


def test_api_rejects_bad_template_ids_versions_and_upload_leaves_no_temp_files(api, tmp_path):
    c, main, service = api
    up = os.path.join(str(tmp_path), "uploads")
    # a garbage SAV + a docx: both temp files must be gone afterwards
    r = c.post("/api/projects", files={"sav": ("a.sav", b"garbage", "application/octet-stream"), "qnr": ("q.docx", b"zzz", "application/octet-stream")})
    assert r.status_code == 400 and os.listdir(up) == []
    r = c.post("/api/projects", files={"sav": ("a.sav", b"garbage", "application/octet-stream")}, data={"dictionary": "nope.xlsx"})
    assert r.status_code in (400, 410) and os.listdir(up) == []
    r = c.post("/api/projects", files={"sav": ("a.sav", b"x", "application/octet-stream"), "qnr": ("q.txt", b"zzz", "application/octet-stream")})
    assert r.status_code == 400 and os.listdir(up) == []                           # wrong second file: the first upload is removed too
    r = c.post("/api/projects/0123456789ab/sav", files={"sav": ("a.sav", b"x", "application/octet-stream")})
    assert r.status_code == 404 and os.listdir(up) == []
    assert c.post("/api/projects/0123456789ab/sav", files={"sav": ("a.txt", b"x", "application/octet-stream")}).status_code in (400, 404)
    assert c.get("/api/docs").status_code in (404, 405) and c.get("/api/openapi.json").status_code in (404, 405)       # no Swagger / CDN


def test_download_names_are_safe_and_gone_errors_do_not_leak_paths(api):
    c, main, service = api
    assert main._safe_name('a/b"c\r\nd') == "a_b_c_d"
    r = c.get("/api/projects/0123456789ab/download")
    assert r.status_code == 404
    from app.runtime import jobs, UserError

    def boom(progress):
        raise FileNotFoundError(2, "No such file", r"C:\secret\projects\x\output.xlsx")
    j = jobs.submit("p", boom)
    for _ in range(100):
        if j["status"] in ("error", "done"):
            break
        import time
        time.sleep(0.05)
    assert j["status"] == "error" and "secret" not in j["error"] and "output.xlsx" not in j["error"]


def test_basic_auth_requires_a_password_and_exact_health_path(api, monkeypatch):
    c, main, service = api
    from app import config
    monkeypatch.setattr(config, "BASIC_AUTH_USER", "u")
    monkeypatch.setattr(config, "BASIC_AUTH_PASS", "")
    import base64
    h = {"authorization": "Basic " + base64.b64encode(b"u:").decode()}
    assert c.get("/api/projects", headers=h).status_code == 401                        # empty password never accepted
    assert c.get("/api/health").status_code in (200, 503)
    assert c.get("/api/healthXYZ").status_code == 401                                  # only the exact path is exempt
    monkeypatch.setattr(config, "BASIC_AUTH_PASS", "סיסמה")
    h = {"authorization": "Basic " + base64.b64encode("u:סיסמה".encode()).decode()}
    assert c.get("/api/projects", headers=h).status_code == 200                       # non-ASCII password works
