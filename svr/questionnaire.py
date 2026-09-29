# -*- coding: utf-8 -*-
"""svr.questionnaire — extract question items from the Word questionnaire (spec 5א).

The questionnaire is used to VALIDATE and to SUGGEST (message order); it is never the primary identification
source and never supplies campaign facts (correct slogan, tested brand, proven exposure).

    parse_docx(path_or_file) -> Questionnaire(items=[Item...], messages=[Message...], warnings=[...])
    cross_validate(qnr, mapping, df, meta) -> list of warnings (Hebrew)
    fingerprint_text(qnr) -> stable string for questionnaire-level template matching
"""
import re
from dataclasses import dataclass, field, asdict

from rapidfuzz import fuzz

TAG_RE = re.compile(r"^(?P<tag>[A-Z][A-Z0-9_\-]{2,}(?: [A-Z][A-Z0-9_\-]{2,})?)\s*[\.:]?\s*(?P<rest>[^A-Za-z]*.*)?$")
NUM_RE = re.compile(r"^\s*\d+\s*[\.:\)]\s*")
UNDERSCORE_RE = re.compile(r"^[_\s\-\.…]+$")
INSTR_WORDS = ("הצג", "תשובה אחת", "אפשר יותר", "ניתן לסמן", "מספר תשובות", "חובה", "לא לשנות", "בסדר משתנה",
               "סיים ר", "הפסק ר", "עבור ל", "רנדומ", "הקפידי", "פתוח")
MSG_ROLE_WORDS = (("מסר עיקרי", "main"), ("מסר משני", "secondary"), ("מסר גינרי", "generic"), ("להסחה", "generic"))
DK_WORDS = ("לא יודע", "אין דעה", "לא בטוח")
NONE_WORDS = ("אף אחד", "לא ראיתי", "אין לי")


@dataclass
class Item:
    tag: str = ""
    question: str = ""
    instructions: list = field(default_factory=list)
    answers: list = field(default_factory=list)
    para: int = 0


@dataclass
class Message:
    text: str
    role: str = ""       # main | secondary | generic | ""  (as annotated in the questionnaire)


@dataclass
class Questionnaire:
    items: list = field(default_factory=list)
    messages: list = field(default_factory=list)      # answer list of the MAIN_MESSAGE_TAKEOUT item, in order
    warnings: list = field(default_factory=list)

    def by_tag(self):
        return {norm_tag(i.tag): i for i in self.items if i.tag}

    def to_dict(self):
        return dict(items=[asdict(i) for i in self.items], messages=[asdict(m) for m in self.messages],
                    warnings=self.warnings)


def norm_tag(t):
    return re.sub(r"[\s_\-]+", "_", str(t)).upper().strip("_")


def _is_instr(t):
    return any(w in t for w in INSTR_WORDS)


def _clean_answer(t):
    t = NUM_RE.sub("", t).strip()
    t = re.sub(r"\s*[-–>]{1,2}\s*(סיים|הפסק|עבור|המשך).*$", "", t).strip()
    t = re.sub(r"\s*>>.*$", "", t).strip()
    return t


def parse_paragraphs(paras):
    """paras: list[str] (already stripped, empty kept as '') -> Questionnaire."""
    q = Questionnaire()
    cur = None
    gap = False
    for i, t in enumerate(paras):
        if not t:
            gap = True
            continue
        was_gap, gap = gap, False
        m = TAG_RE.match(t)
        is_tag = bool(m) and not re.search(r"[֐-׿]", m.group("tag")) and len(t) < 90
        if is_tag:
            cur = Item(tag=m.group("tag").strip(), para=i)
            rest = (m.group("rest") or "").strip()
            if rest:
                cur.instructions.append(rest)
            q.items.append(cur)
            continue
        if UNDERSCORE_RE.match(t):
            continue
        # a blank line after a finished answer list + a new question-like line = a new (untagged) item
        if cur is not None and was_gap and len(cur.answers) >= 2 and (t.endswith("?") or t.endswith(":")) and len(t) > 25:
            cur = Item(tag="", question=t, para=i)
            q.items.append(cur)
            continue
        if cur is None:
            continue
        if _is_instr(t) and len(t) < 70:
            cur.instructions.append(t)
        elif not cur.question and (t.endswith("?") or t.endswith(":") or len(t) > 40):
            cur.question = t
        elif cur.question and (t.endswith("?") and len(t) > 30):
            cur.question += " " + t
        else:
            if len(t) <= 110 and ".mp4" not in t:
                cur.answers.append(_clean_answer(t))
    # message roles (annotation lives on the answer line or on the next short line "- מסר עיקרי")
    for it in q.items:
        if norm_tag(it.tag) == "MAIN_MESSAGE_TAKEOUT":
            msgs = []
            for a in it.answers:
                role = next((r for w, r in MSG_ROLE_WORDS if w in a), "")
                if a.strip(" -–") in [w for w, _ in MSG_ROLE_WORDS] and msgs:
                    msgs[-1].role = role or msgs[-1].role
                    continue
                if any(w in a for w in DK_WORDS + NONE_WORDS):
                    msgs.append(Message(a, "none"))
                    continue
                clean = a
                for w, _ in MSG_ROLE_WORDS:
                    clean = clean.replace(w, "")
                msgs.append(Message(clean.strip(" -–"), role))
            q.messages = msgs
    if not q.items:
        q.warnings.append("לא זוהו תגיות שאלות בשאלון (שורות באותיות גדולות) — האימות מול השאלון לא בוצע")
    if not q.messages:
        q.warnings.append("לא נמצאה רשימת מסרים (MAIN_MESSAGE_TAKEOUT) בשאלון — סדר המסרים יילקח לפי מיקום")
    return q


def parse_docx(path_or_file):
    """Best effort: a failed parse returns an empty Questionnaire with a warning (the app continues without it)."""
    try:
        import docx
        d = docx.Document(path_or_file)
        paras = [p.text.strip() for p in d.paragraphs]
        q = parse_paragraphs(paras)
        # table cells (e.g. 'AD_ RELEVANCE' statements) are appended as extra items
        for t in d.tables:
            first = " ".join(c.text.strip() for c in t.rows[0].cells if c.text.strip())
            m = TAG_RE.match(first.replace(" ", "_", 1)) if first else None
            if m and not re.search(r"[֐-׿]", m.group("tag")):
                it = Item(tag=m.group("tag"), question="", para=-1)
                for r in t.rows[1:]:
                    txt = " ".join(c.text.strip() for c in r.cells if c.text.strip())
                    if txt:
                        it.answers.append(txt[:200])
                q.items.append(it)
        return q
    except Exception as ex:  # noqa: BLE001 - never block the run
        q = Questionnaire()
        q.warnings.append(f"פענוח השאלון נכשל ({type(ex).__name__}) — ממשיכים בלעדיו")
        return q


# ---------------------------------------------------------------------------------------------
def match_item(qnr, key, dict_var, question):
    """Find the questionnaire item for a SAV block: tag equality first, then question-text similarity."""
    tags = qnr.by_tag()
    for cand in (key, dict_var or ""):
        n = norm_tag(cand)
        if n and n in tags:
            return tags[n], "tag"
    best, bs = None, 0
    qq = re.sub(r"\[[^\]]*\]", " ", question or "")
    if len(qq.strip()) < 20:
        return None, ""
    for it in qnr.items:
        if not it.question:
            continue
        s = fuzz.token_set_ratio(qq[:140], it.question[:140])
        if s > bs:
            best, bs = it, s
    return (best, "text") if best and bs >= 80 else (None, "")


def _sav_options(e, df, meta):
    if e.get("option_names"):
        return list(e["option_names"].values())
    if e.get("code_names"):
        return list(e["code_names"].values())
    if len(e["vars"]) == 1:
        vl = (meta.variable_value_labels or {}).get(e["vars"][0]) or {}
        return [str(v) for v in vl.values()]
    return []


NOISE_RE = re.compile(r"(הצג|יוצג|צייני|צייני|בחר|כעת|אנא|חשיפה ל|נא |סמני|לצרף|לעדכן|^- |מסר (עיקרי|משני)|שאל[:.]?$|[:\-–]$)")


def _plausible_answer(a):
    return 0 < len(a) <= 70 and not NOISE_RE.search(a)


def cross_validate(qnr, mapping, df, meta, threshold=70):
    """Hebrew warnings: answers that exist on one side only.  Also stamps `qnr_item` on matched entries."""
    out = []
    for e in mapping.get("questions", []):
        it, how = match_item(qnr, e["key"], e.get("dict_var"), e.get("question", ""))
        if not it or not it.answers:
            continue
        e["qnr_item"] = dict(tag=it.tag, matched_by=how)
        if how != "tag" or e.get("type") in ("coded_open", "describe", "message_takeout") or "MESSAGE_TAKEOUT" in norm_tag(it.tag):
            continue          # text matches are unreliable and coded lists are researcher-defined: no warnings from them
        sav = [re.sub(r"\[[^\]]*\]", "", s).strip() for s in _sav_options(e, df, meta)]
        sav = [s for s in sav if s]
        if not sav:
            continue
        ans = [re.sub(r"\[[^\]]*\]", "", a).strip() for a in it.answers if a.strip()]
        ans = [a for a in ans if _plausible_answer(a)]
        if not ans:
            continue
        for a in ans:
            if not any(fuzz.token_set_ratio(a, s) >= threshold for s in sav):
                out.append(f"{e['key']}: התשובה \"{a[:50]}\" קיימת בשאלון ולא ב-SAV")
        for s in sav:
            if not any(fuzz.token_set_ratio(a, s) >= threshold for a in ans):
                out.append(f"{e['key']}: התשובה \"{s[:50]}\" קיימת ב-SAV ולא בשאלון")
        if e.get("type") in ("scale", "single") and len(sav) != len(ans) and how == "tag":
            out.append(f"{e['key']}: מספר התשובות ב-SAV ({len(sav)}) שונה מהשאלון ({len(ans)})")
    return out


def suggest_message_order(qnr, mapping):
    """Suggested position of each message option in MAIN/TOTAL_MESSAGE_TAKEOUT — from the roles annotated in the
    questionnaire.  Returns dict(order=[...], source='questionnaire') or None.  A suggestion only (needs approval)."""
    real = [m for m in qnr.messages if m.role != "none"]
    if not real:
        return None
    roles = [m.role or "?" for m in real]
    return dict(order=[dict(text=m.text[:120], role=m.role or "") for m in real], roles=roles,
                source="questionnaire", note="הסדר והתפקידים נלקחו מהשאלון — דורש אישור")


def fingerprint_text(qnr):
    return "|".join(sorted(norm_tag(i.tag) for i in qnr.items if i.tag))
