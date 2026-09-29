# -*- coding: utf-8 -*-
"""svr_lib — shared helpers for the Impact360 SAV runner.

Loads the bench dictionary, reads the SAV, detects question blocks, and holds the
statistical primitives (weighted %, two-proportion z-test).  No Excel code here.
"""
import glob
import os
import re
from collections import OrderedDict

import numpy as np
import openpyxl
import pandas as pd
import pyreadstat

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_ROOT = os.path.dirname(HERE)
BUNDLED_DICT = os.path.join(SKILL_ROOT, "assets", "Impact360_import_dictionary.xlsx")

# ----------------------------------------------------------------------------
# Dictionary
# ----------------------------------------------------------------------------

def find_dictionary(explicit=None, search_roots=None):
    """Newest 'Impact360 import dictionary*.xlsx' (not a 'Copy of', not in OLD) or the bundled one."""
    if explicit and os.path.exists(explicit):
        return explicit
    cands = []
    for root in (search_roots or []):
        for p in glob.glob(os.path.join(root, "**", "*import dictionary*.xlsx"), recursive=True):
            low = p.lower()
            if "\\old\\" in low or "/old/" in low or os.path.basename(low).startswith(("copy of", "~$")):
                continue
            cands.append(p)
    if cands:
        return max(cands, key=os.path.getmtime)
    return BUNDLED_DICT


def _num(x, default=999):
    try:
        return float(x)
    except (TypeError, ValueError):
        return default


def load_dictionary(path):
    """Return (vars, codes).
    vars[VAR] = dict(title, question, module, type, answers, base, status)
    codes[VAR] = list of dict(slot, role, label, united, status, polarity)
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    vars_ = OrderedDict()
    for r in wb["DATA_Dictionary"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        vars_[str(r[0])] = dict(title=r[1] or "", question=(r[2] or ""), module=r[3] or "", type=r[4] or "",
                                answers=r[5] or "", base=r[6] or "", status=r[8] or "",
                                chapter=_num(r[9]), varsort=_num(r[10]))
    codes = OrderedDict()
    for r in wb["DATA_Dictionary_CODES"].iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        codes.setdefault(str(r[0]), []).append(dict(
            slot=int(r[4]) if str(r[4]).strip().lstrip("-").isdigit() else r[4], role=r[5] or "",
            label=r[6] or "", united=r[9] or "", status=r[10] or "", polarity=r[11] or ""))
    return vars_, codes


def base_kind(base_text):
    b = (base_text or "").lower()
    if "exposed" in b or "נחשפים" in b:
        return "exposed"
    if "computed" in b or "מחושב" in b:
        return "computed"
    return "sample"


# ----------------------------------------------------------------------------
# SAV
# ----------------------------------------------------------------------------

def load_sav(path):
    """Data + metadata.  SPSS value labels written by Hebrew Decipher exports come out scrambled with the
    default cp1255 decode (a stray first letter, last letter lost); the same bytes read as iso-8859-8 are clean,
    so value labels are re-read that way (variable labels are fine either way)."""
    df, meta = pyreadstat.read_sav(path)
    try:
        _, m2 = pyreadstat.read_sav(path, encoding="iso-8859-8", metadataonly=True)
        if m2.variable_value_labels:
            meta.variable_value_labels = m2.variable_value_labels
            meta.value_labels = getattr(m2, "value_labels", meta.value_labels)
    except Exception:
        pass
    return df, meta


_NAME_PREFIX = re.compile(r"^.{0,1}?(?=[A-Za-z_][\w\-\.]*:\s)")


def split_label(name, label):
    """'q6r1: האגיס[huggies.jpg] - האם ...' -> (option, question).  Decipher grid labels are 'option - question'."""
    lab = str(label or "")
    m = re.match(r"^.?" + re.escape(name) + r":\s*(.*)$", lab, flags=re.S)
    body = m.group(1) if m else lab
    if " - " in body:
        opt, q = body.split(" - ", 1)
        return opt.strip(), q.strip()
    return "", body.strip()


BLOCK_RE = re.compile(r"^(?P<stem>.+?)(?:r\d+c\d+|r\d+|c\d+|none|oe)$")


def block_stem(name):
    m = BLOCK_RE.match(name)
    return m.group("stem") if m else name


def detect_blocks(df, meta):
    """Group SAV columns into question blocks by Decipher naming (stem + rN/cN)."""
    blocks = OrderedDict()
    for n in meta.column_names:
        blocks.setdefault(block_stem(n), []).append(n)
    return blocks


def is_binary(series):
    """0/1 flag column.  An all-empty column (option never shown) counts as binary."""
    vals = set(series.dropna().unique().tolist())
    return vals <= {0, 1, 0.0, 1.0}


# ----------------------------------------------------------------------------
# Statistics
# ----------------------------------------------------------------------------

def wpct(cond, valid, w=None):
    """Weighted % of `cond` among `valid`.  Both boolean arrays.  Returns (pct, n_unweighted)."""
    n = int(valid.sum())
    if n == 0:
        return None, 0
    if w is None:
        return 100.0 * float(cond[valid].sum()) / n, n
    ws = float(w[valid].sum())
    if ws == 0:
        return None, n
    return 100.0 * float(w[valid & cond].sum()) / ws, n


def wmean(values, valid, w=None):
    n = int(valid.sum())
    if n == 0:
        return None, 0
    if w is None:
        return float(values[valid].mean()), n
    ws = float(w[valid].sum())
    return float((values[valid] * w[valid]).sum() / ws), n


def z_two_prop(p1, n1, p2, n2, min_n=30):
    """+1 / -1 / 0 : is p1 significantly higher / lower / not different from p2 (95%, pooled)."""
    if p1 is None or p2 is None or n1 < min_n or n2 < min_n:
        return 0
    p1, p2 = p1 / 100.0, p2 / 100.0
    pp = (p1 * n1 + p2 * n2) / (n1 + n2)
    se = np.sqrt(pp * (1 - pp) * (1 / n1 + 1 / n2))
    if se == 0:
        return 0
    z = (p1 - p2) / se
    return 1 if z > 1.96 else (-1 if z < -1.96 else 0)


def z_two_mean(m1, s1, n1, m2, s2, n2, min_n=30):
    if None in (m1, m2, s1, s2) or n1 < min_n or n2 < min_n:
        return 0
    se = np.sqrt(s1 ** 2 / n1 + s2 ** 2 / n2)
    if se == 0:
        return 0
    z = (m1 - m2) / se
    return 1 if z > 1.96 else (-1 if z < -1.96 else 0)


# ----------------------------------------------------------------------------
# DESCRIBE worlds — Bible v5.2, sheet DESCRIBE_Worlds_Pentagon (W-codes by attribute name)
# ----------------------------------------------------------------------------
DESCRIBE_WORLDS = {
    # attribute (Hebrew, as in dictionary/questionnaire) : world
    "מושכת תשומת לב": "Magnetism", "יצירתית": "Magnetism", "מקורית": "Magnetism", "מצחיקה": "Magnetism",
    "זכירה": "Magnetism", "חזקה": "Magnetism", "ייחודית": "Magnetism",
    "מרגשת": "Emotional", "מחממת את הלב": "Emotional",
    "אקטיבית": "Dynamic", "שמחה": "Dynamic", "אנרגטית": "Dynamic",
    "מעצבנת": "Negative", "מעליבה": "Negative",
    "אמינה": "Rational", "אינפורמטיבית": "Rational", "מציאותית": "Rational", "עושה את העבודה": "Rational",
    "משעממת": "-", "נעימה": "-",
}
WORLD_SLOT = {"Magnetism": 21, "Emotional": 22, "Dynamic": 23, "Negative": 24, "Rational": 25}


def z_prop(p1, n1, p2, n2, min_n=30):
    """signed z of p1 - p2 (pooled two-proportion); 0 when a base is below min_n"""
    if p1 is None or p2 is None or n1 < min_n or n2 < min_n:
        return 0.0
    a, b = p1 / 100.0, p2 / 100.0
    pp = (a * n1 + b * n2) / (n1 + n2)
    se = np.sqrt(pp * (1 - pp) * (1 / n1 + 1 / n2))
    return 0.0 if se == 0 else float((a - b) / se)


def z_mean(m1, s1, n1, m2, s2, n2, min_n=30):
    if None in (m1, m2, s1, s2) or n1 < min_n or n2 < min_n:
        return 0.0
    se = np.sqrt(s1 ** 2 / n1 + s2 ** 2 / n2)
    return 0.0 if se == 0 else float((m1 - m2) / se)
